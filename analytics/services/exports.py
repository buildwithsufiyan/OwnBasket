import csv
from io import BytesIO

from django.http import HttpResponse, StreamingHttpResponse
from django.utils import timezone


class Echo:
    def write(self, value):
        return value


def safe_cell(value):
    if value is None:
        return ''
    if hasattr(value, 'isoformat'):
        value = timezone.localtime(value).isoformat() if hasattr(value, 'tzinfo') and value.tzinfo else value.isoformat()
    value = str(value)
    if value.startswith(('=', '+', '-', '@', '\t', '\r')):
        return "'" + value
    return value


def csv_response(filename, headers, rows):
    writer = csv.writer(Echo())

    def stream():
        yield '\ufeff'
        yield writer.writerow(headers)
        for row in rows:
            yield writer.writerow([safe_cell(value) for value in row])

    response = StreamingHttpResponse(stream(), content_type='text/csv; charset=utf-8')
    response['Content-Disposition'] = f'attachment; filename="{filename}.csv"'
    response['Cache-Control'] = 'no-store'
    return response


def excel_response(filename, title, headers, rows):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        return HttpResponse('Excel export requires the openpyxl production dependency.', status=503, content_type='text/plain')

    workbook = Workbook(write_only=False)
    summary = workbook.active
    summary.title = 'Summary'
    summary.append([title])
    summary.append(['Generated', timezone.localtime().isoformat()])
    summary.append(['Rows', 'See Data sheet'])
    data = workbook.create_sheet('Data')
    data.append(headers)
    for cell in data[1]:
        cell.font = Font(bold=True, color='FFFFFF')
        cell.fill = PatternFill('solid', fgColor='232F3E')
    count = 0
    for row in rows:
        data.append([safe_cell(value) for value in row])
        count += 1
    summary['B3'] = count
    data.freeze_panes = 'A2'
    for index, header in enumerate(headers, 1):
        data.column_dimensions[get_column_letter(index)].width = min(max(len(str(header)) + 2, 12), 32)
    output = BytesIO()
    workbook.save(output)
    response = HttpResponse(output.getvalue(), content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="{filename}.xlsx"'
    response['Cache-Control'] = 'no-store'
    return response
